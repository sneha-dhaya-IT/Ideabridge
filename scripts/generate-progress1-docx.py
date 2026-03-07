"""
Generate the Progress 1 — UI & Business Rules Word document
with all screenshots embedded.

Run:  python scripts/generate-progress1-docx.py
"""

import os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCREENSHOTS = os.path.join(BASE, "screenshots")
OUT = os.path.join(BASE, "docs", "Progress-1-UI-Business-Rules.docx")


def set_cell_shading(cell, color_hex):
    """Set background color of a table cell."""
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(
        qn("w:shd"),
        {
            qn("w:val"): "clear",
            qn("w:color"): "auto",
            qn("w:fill"): color_hex,
        },
    )
    shading.append(shading_elem)


def style_header_row(row, bg="1E3A5F"):
    """Style a table header row with dark blue background and white bold text."""
    for cell in row.cells:
        set_cell_shading(cell, bg)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.bold = True
                run.font.size = Pt(9)


def add_table(doc, headers, rows, col_widths=None):
    """Create a styled table with header row and data rows."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_header_row(table.rows[0])

    # Data
    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)

    # Column widths
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)

    doc.add_paragraph()  # spacer
    return table


def add_screenshot(doc, filename, caption, width=5.8):
    """Add a screenshot image with a caption."""
    path = os.path.join(SCREENSHOTS, filename)
    if not os.path.exists(path):
        doc.add_paragraph(f"[Screenshot not found: {filename}]")
        return
    doc.add_picture(path, width=Inches(width))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.font.size = Pt(8)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    return h


def main():
    doc = Document()

    # ── Page margins ──
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # ══════════════════════════════════════════════════════════════
    # COVER / TITLE
    # ══════════════════════════════════════════════════════════════
    doc.add_paragraph()
    title = doc.add_heading("Progress 1", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
        run.font.size = Pt(28)

    sub = doc.add_heading("UI and Business Rules Document", level=1)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in sub.runs:
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
        run.font.size = Pt(16)

    doc.add_paragraph()

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in [
        "Module: IT3040 – ITPM | Semester 1",
        "Programme: BSc (Hons) in Information Technology — Year 3",
        "Date: March 2026",
    ]:
        run = meta.add_run(line + "\n")
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # 1. STUDENT DETAILS
    # ══════════════════════════════════════════════════════════════
    heading(doc, "1. Student Details")

    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Student Name", "Akshayan"],
            ["Registration Number", "(fill in your SLIIT reg. no.)"],
            ["Responsible Component(s)", "Idea & Guidance Module (Member 2)"],
            [
                "Description",
                "A web module that lets students submit project ideas with "
                "validation, receive recursive threaded feedback from mentors "
                "and peers, and view guidance threads per project.",
            ],
        ],
        col_widths=[5, 12],
    )

    # ══════════════════════════════════════════════════════════════
    # 2. COMPONENT SUMMARY
    # ══════════════════════════════════════════════════════════════
    heading(doc, "2. Component Summary")

    add_table(
        doc,
        ["#", "UI Screen", "Route", "Key Screenshot"],
        [
            ["UI 1", "Project Idea Submission Form", "/posts/new", "01-post-form-empty.png"],
            ["UI 2", "Form Validation Error View", "/posts/new (invalid)", "02-post-form-validation-errors.png"],
            ["UI 3", "Feedback Thread (Comment Tree)", "/feedback", "05-feedback-thread-full.png"],
            ["UI 4", "Guidance Thread", "/guidance/[projectId]", "09-guidance-thread.png"],
        ],
        col_widths=[1.5, 5, 4, 6.5],
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # 3. UI 1 — POST FORM
    # ══════════════════════════════════════════════════════════════
    heading(doc, "3. UI 1: Project Idea Submission Form (/posts/new)")

    add_screenshot(doc, "01-post-form-empty.png", "Figure 1 — PostForm empty state")
    add_screenshot(doc, "03-post-form-filled.png", "Figure 2 — PostForm filled with ITPM test data")

    # A) Access Control
    heading(doc, "A) Access Control Rules", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["AC-1", "Only authenticated (logged-in) students can submit a project idea. Unauthenticated users receive an error on submission."],
            ["AC-2", "The form page itself is publicly viewable, but the submit action requires a valid Supabase session."],
        ],
        col_widths=[1.5, 15.5],
    )

    # B) Validation Rules
    heading(doc, "B) Validation Rules (Input Constraints)", level=2)
    add_table(
        doc,
        ["#", "Field", "Rule"],
        [
            ["VR-1", "Title", "Mandatory. Minimum 10 characters. Must be a non-empty string."],
            ["VR-2", "Problem Statement", "Mandatory. Must be a non-empty string."],
            ["VR-3", "Project Variant", 'Mandatory. Must be one of: research, prototype, capstone, mini-project.'],
            ["VR-4", "Tech Stacks", "Mandatory. At least one tag must be selected. Allowed values: Next.js, React, TypeScript, Tailwind CSS, Node.js, Python, Supabase, PostgreSQL, Prisma, Docker."],
            ["VR-5", "URLs", "Optional. Each entry must be a valid URL format."],
            ["VR-6", "All fields", "Validation is performed server-side using Zod schema. Client-side feedback is displayed inline under each invalid field."],
        ],
        col_widths=[1.5, 3.5, 12],
    )

    # C) Process / Workflow
    heading(doc, "C) Process / Workflow Rules", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["WF-1", 'User fills in form fields → clicks "Submit Idea" → server action validates with Zod → if valid, inserts row into project_ideas table in Supabase.'],
            ["WF-2", "On successful submission, the home page (/) is revalidated and the user sees a success confirmation."],
            ["WF-3", "On validation failure, the form re-renders with inline error messages below each invalid field. Data already entered is preserved."],
            ["WF-4", 'The submit button shows "Submitting…" and is disabled while the server action is in progress, preventing double submissions.'],
        ],
        col_widths=[1.5, 15.5],
    )

    # D) Data Consistency
    heading(doc, "D) Data Consistency Rules", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["DC-1", "Each submitted idea is stored with the authenticated user's ID (user_id from Supabase auth)."],
            ["DC-2", "The tech_stack field is stored as a JSON array in the database."],
            ["DC-3", "The variant field is stored as one of the four allowed enum values."],
        ],
        col_widths=[1.5, 15.5],
    )

    # E) System Response
    heading(doc, "E) Notification / System Response Rules", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["NR-1", "On successful insert, a success message is returned and the page is revalidated."],
            ["NR-2", "On validation failure, individual field errors are displayed in red text beneath the respective input."],
            ["NR-3", "On server/database error, a generic error message is displayed to the user."],
        ],
        col_widths=[1.5, 15.5],
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # 4. UI 2 — VALIDATION ERROR VIEW
    # ══════════════════════════════════════════════════════════════
    heading(doc, "4. UI 2: Form Validation Error View (/posts/new — invalid submit)")

    add_screenshot(doc, "02-post-form-validation-errors.png", "Figure 3 — All validation errors displayed")
    add_screenshot(doc, "04-post-form-title-error.png", "Figure 4 — Title minimum-length error")

    heading(doc, "B) Validation Rules (Input Constraints)", level=2)
    add_table(
        doc,
        ["#", "Scenario", "Expected Behavior"],
        [
            ["VE-1", "All fields empty → Submit", 'Errors shown for: Title ("at least 10 characters"), Problem Statement ("Required"), Tech Stacks ("Select at least one").'],
            ["VE-2", "Title fewer than 10 characters", '"Title must be at least 10 characters."'],
            ["VE-3", "No variant selected", '"Please select a project variant."'],
            ["VE-4", "No tech stacks selected", '"Select at least one technology."'],
            ["VE-5", "Invalid URL format entered", "URL-specific validation message."],
        ],
        col_widths=[1.5, 5, 10.5],
    )

    heading(doc, "F) System Response Rules", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["SR-1", "Error messages appear inline in red (text-red-600) directly below each invalid field."],
            ["SR-2", "The form is not submitted when validation fails — no database write occurs."],
            ["SR-3", "Previously entered valid data is preserved when the form re-renders with errors."],
        ],
        col_widths=[1.5, 15.5],
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # 5. UI 3 — FEEDBACK THREAD
    # ══════════════════════════════════════════════════════════════
    heading(doc, "5. UI 3: Feedback Thread — Comment Tree (/feedback)")

    add_screenshot(doc, "05-feedback-thread-full.png", "Figure 5 — Full feedback thread with nested comments")
    add_screenshot(doc, "06-feedback-upvoted.png", "Figure 6 — Upvoted comment state")

    doc.add_page_break()

    add_screenshot(doc, "07-feedback-accepted.png", "Figure 7 — Accepted comment state (OP view)")
    add_screenshot(doc, "08-feedback-upvote-and-accept.png", "Figure 8 — Upvote + Accept active simultaneously")

    # A) Access Control
    heading(doc, "A) Access Control Rules", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["AC-3", "The feedback thread page is viewable by all users (public)."],
            ["AC-4", 'The "Mark Accepted" button is visible only to the Original Poster (OP). Other users cannot see or interact with it.'],
            ["AC-5", 'The "Upvote" button is available to all users viewing the thread.'],
        ],
        col_widths=[1.5, 15.5],
    )

    # B) Display
    heading(doc, "B) Display Rules", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["DR-1", "Comments are rendered as a recursive tree. Root comments appear at the top level; replies are indented with a left blue border."],
            ["DR-2", "Nesting depth is unlimited — replies to replies render at increasing indentation."],
            ["DR-3", 'Mentor comments display an amber badge (text "Mentor") next to the author name.'],
            ["DR-4", 'OP (Original Poster) comments display a blue badge (text "OP") next to the author name.'],
            ["DR-5", "Student comments have no special badge."],
            ["DR-6", "Comment body supports full Markdown: bold, links, inline code, and fenced code blocks with syntax highlighting."],
        ],
        col_widths=[1.5, 15.5],
    )

    # C) Upvote
    heading(doc, "C) Interaction Rules — Upvote Toggle", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["UV-1", 'Each comment has an "Upvote" button. Clicking it toggles the state to "Upvoted".'],
            ["UV-2", 'When upvoted, the button text changes to "Upvoted", aria-pressed changes from false to true, and the button background changes to blue.'],
            ["UV-3", 'Clicking "Upvoted" toggles back to the initial "Upvote" state (aria-pressed="false", default styling).'],
            ["UV-4", "Upvote state is maintained per-comment and is independent of other comments."],
        ],
        col_widths=[1.5, 15.5],
    )

    # D) Accept
    heading(doc, "D) Interaction Rules — Accept Solution Toggle (OP Only)", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["AS-1", 'The "Mark Accepted" button appears only when isOP is true. It is rendered on every comment in the thread.'],
            ["AS-2", 'Clicking "Mark Accepted" toggles the state to "Accepted" with aria-pressed="true" and a green background.'],
            ["AS-3", 'Clicking "Accepted" toggles back to "Mark Accepted" (aria-pressed="false", default styling).'],
            ["AS-4", "Both Upvote and Accept can be active simultaneously on the same comment. They are independent toggles."],
        ],
        col_widths=[1.5, 15.5],
    )

    # E) Data Consistency
    heading(doc, "E) Data Consistency Rules", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["DC-4", "Upvote and accept states are currently client-side only (React useState). They reset on page reload."],
            ["DC-5", "Comment data (author, role, content, parent_id) is passed as props from the server component."],
        ],
        col_widths=[1.5, 15.5],
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # 6. UI 4 — GUIDANCE THREAD
    # ══════════════════════════════════════════════════════════════
    heading(doc, "6. UI 4: Guidance Thread (/guidance/[projectId])")

    add_screenshot(doc, "09-guidance-thread.png", "Figure 9 — Guidance Thread page (empty state)")

    # A) Access
    heading(doc, "A) Access Control Rules", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["AC-6", "The guidance thread page is accessible via dynamic route /guidance/[projectId]. Any user with the URL can view it."],
            ["AC-7", "Comments are fetched from the Supabase comments table filtered by project_id."],
        ],
        col_widths=[1.5, 15.5],
    )

    # B) Display
    heading(doc, "B) Display Rules", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["DR-7", 'The page heading reads "Guidance Thread" with the projectId displayed in a code tag in the subtitle.'],
            ["DR-8", 'If no comments exist for the given projectId, an empty state message is shown: "No guidance comments yet — be the first to reply!"'],
            ["DR-9", "If comments exist, they are rendered as a recursive tree identical to the Feedback Thread structure (nested, indented, with role badges)."],
            ["DR-10", "Each comment shows a created_at timestamp in addition to author and content."],
        ],
        col_widths=[1.5, 15.5],
    )

    # C) Data Fetch
    heading(doc, "C) Data Fetch Rules", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["DF-1", "Comments are fetched server-side (async Server Component) using fetchCommentsByProjectId()."],
            ["DF-2", "Comments are ordered by created_at ascending to ensure parent nodes are processed before children."],
            ["DF-3", "If Supabase environment variables are not configured, the function returns an empty array gracefully (no crash)."],
            ["DF-4", "If the Supabase query fails, the error is logged to the server console and an empty array is returned."],
        ],
        col_widths=[1.5, 15.5],
    )

    # D) Tree Building
    heading(doc, "D) Tree Building Rules", level=2)
    add_table(
        doc,
        ["#", "Rule"],
        [
            ["TB-1", "The buildCommentTree() function converts a flat array of comments into a nested tree using an O(n) single-pass algorithm."],
            ["TB-2", "Comments with parent_id: null become root nodes. Comments with a valid parent_id are nested under their parent's children array."],
            ["TB-3", "Orphan comments (with a parent_id that doesn't match any existing comment) are promoted to root level."],
        ],
        col_widths=[1.5, 15.5],
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════
    # 7. TESTING EVIDENCE
    # ══════════════════════════════════════════════════════════════
    heading(doc, "7. Testing Evidence")

    heading(doc, "Unit Tests (Vitest)", level=2)
    p = doc.add_paragraph()
    run = p.add_run("7 tests")
    run.bold = True
    p.add_run(" covering the createProjectIdea server action: validation of all fields, "
              "Zod schema enforcement, Supabase integration error handling.")
    doc.add_paragraph()

    heading(doc, "E2E Tests (Playwright)", level=2)
    p = doc.add_paragraph()
    run = p.add_run("25 tests")
    run.bold = True
    p.add_run(" across 8 describe blocks covering all 4 UIs:")

    tests = [
        "PostForm field validation (3 tests)",
        "PostForm submission with ITPM test data (2 tests)",
        "FeedbackThread comment tree rendering (4 tests)",
        "FeedbackThread Markdown rendering (4 tests)",
        "FeedbackThread role badges (3 tests)",
        "FeedbackThread upvote interaction (2 tests)",
        "FeedbackThread accept solution toggle (4 tests)",
        "GuidanceThread page rendering (3 tests)",
    ]
    for t in tests:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_paragraph()

    heading(doc, "Version Control (Git)", level=2)
    git_items = [
        "Repository: https://github.com/AK29-Shay/project",
        "Branch: feature/member2-idea-guidance",
        "Meaningful commits tracking each feature addition",
    ]
    for item in git_items:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph()

    # ══════════════════════════════════════════════════════════════
    # 8. TECHNOLOGY STACK
    # ══════════════════════════════════════════════════════════════
    heading(doc, "8. Technology Stack")

    add_table(
        doc,
        ["Technology", "Purpose"],
        [
            ["Next.js 14 (App Router)", "Full-stack React framework"],
            ["React 18", "UI component library"],
            ["TypeScript 5", "Type-safe development"],
            ["Tailwind CSS 4", "Utility-first styling"],
            ["Zod 4", "Schema validation (server-side)"],
            ["Supabase (SSR)", "Database + Authentication"],
            ["Playwright", "E2E automated testing"],
            ["Vitest", "Unit testing"],
            ["Vercel", "Cloud deployment"],
        ],
        col_widths=[5, 12],
    )

    # ══════════════════════════════════════════════════════════════
    # 9. DEPLOYMENT
    # ══════════════════════════════════════════════════════════════
    heading(doc, "9. Deployment")

    deploy_items = [
        ("Platform:", " Vercel"),
        ("GitHub Integration:", " Auto-deploys from feature/member2-idea-guidance branch"),
        ("Live URL:", " https://project-fawn-six-84.vercel.app"),
    ]
    for label, value in deploy_items:
        p = doc.add_paragraph()
        run = p.add_run(label)
        run.bold = True
        p.add_run(value)

    # ── Save ──
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print(f"Word document saved to: {OUT}")
    print(f"File size: {os.path.getsize(OUT) / 1024:.1f} KB")


if __name__ == "__main__":
    main()
